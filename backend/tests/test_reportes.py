# -*- coding: utf-8 -*-
"""Pruebas de dashboards, exportación, perfil histórico e IA (RF-08, RF-10..RF-14)."""
import csv
import io
import re
from datetime import datetime, timezone
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services import reporte_service
from tests.conftest import headers


def test_dashboard(client, tokens):
    r = client.get("/api/reportes/dashboard", params={"periodo": "2026-S1"},
                   headers=headers(tokens["consulta"]))
    assert r.status_code == 200
    cuerpo = r.json()
    kpis = cuerpo["kpis"]
    assert kpis["evaluaciones_periodo"] > 0
    assert kpis["colaboradores_activos"] >= 12
    assert sum(cuerpo["distribucion_clasificacion"].values()) == kpis["evaluaciones_periodo"]
    assert cuerpo["evolucion_periodos"]


def test_export_csv_respeta_permisos(client, tokens):
    # CONSULTA no tiene permiso de exportar (RN-07 / RF-14)
    r = client.get("/api/reportes/export/evaluaciones.csv",
                   headers=headers(tokens["consulta"]))
    assert r.status_code == 403

    r2 = client.get("/api/reportes/export/evaluaciones.csv",
                    headers=headers(tokens["gestor"]))
    assert r2.status_code == 200
    assert "text/csv" in r2.headers["content-type"]
    assert r2.content.startswith(b"\xef\xbb\xbf")
    filas = list(csv.reader(io.StringIO(r2.content.decode("utf-8-sig"))))
    assert filas[0][0] == "evaluacion_id"
    assert len(filas[0]) == 16
    assert filas[1:]
    assert all(len(fila) == len(filas[0]) for fila in filas[1:])
    assert any(fila[1].startswith("COL-0") for fila in filas[1:])
    assert "\ufffd" not in r2.content.decode("utf-8-sig")
    assert re.fullmatch(r"\d{2}/\d{2}/\d{4} \d{2}:\d{2}:\d{2}", filas[1][6])


def test_export_excel_formateado_respeta_permisos_y_periodo(client, tokens):
    r = client.get("/api/reportes/export/evaluaciones.xlsx",
                   headers=headers(tokens["consulta"]))
    assert r.status_code == 403

    respuesta = client.get(
        "/api/reportes/export/evaluaciones.xlsx",
        params={"periodo": "2026-S1"},
        headers=headers(tokens["gestor"]),
    )
    assert respuesta.status_code == 200
    assert "spreadsheetml.sheet" in respuesta.headers["content-type"]
    assert 'filename="evaluaciones_2026-S1_' in respuesta.headers["content-disposition"]

    libro = load_workbook(io.BytesIO(respuesta.content), data_only=True)
    assert libro.sheetnames == ["Resultados"]
    hoja = libro["Resultados"]
    assert hoja["A1"].value == "RESULTADOS DE EVALUACIONES"
    assert hoja["A2"].value.startswith("Sistema HiPo · Periodo: 2026-S1")
    assert hoja["A2"].value.endswith("hora CDMX")
    assert [hoja.cell(4, columna).value for columna in range(1, 17)] == [
        "evaluacion_id", "colaborador_codigo", "colaborador_nombre", "area", "puesto",
        "periodo", "fecha_evaluacion", "version_modelo", "configuracion_version",
        "estado_evaluacion", "puntuacion", "clasificacion", "cobertura",
        "factores_principales", "estado_revision", "explicacion_origen",
    ]
    assert hoja["A4"].fill.fgColor.rgb.endswith("1F4E78")
    assert hoja.freeze_panes == "A5"
    assert hoja.auto_filter.ref == f"A4:P{hoja.max_row}"
    assert hoja.max_row > 4
    assert all(hoja.cell(fila, 6).value == "2026-S1" for fila in range(5, hoja.max_row + 1))
    libro.close()


def test_export_excel_convierte_fecha_a_hora_cdmx(monkeypatch):
    evaluacion = SimpleNamespace(
        id=1,
        periodo="2026-S1",
        fecha=datetime(2026, 10, 16, 14, 30, tzinfo=timezone.utc),
        version_modelo="reglas-v1",
        configuracion=None,
        estado=SimpleNamespace(value="completada"),
    )
    resultado = SimpleNamespace(
        factores=[],
        revisiones=[],
        puntuacion=84.5,
        clasificacion=SimpleNamespace(value="alto_potencial"),
        cobertura=0.9,
        explicacion_origen="local_determinista",
    )
    colaborador = SimpleNamespace(
        codigo="COL-001", nombre="Colaborador de prueba", area="Tecnología", puesto="Analista"
    )
    monkeypatch.setattr(
        reporte_service,
        "_evaluaciones_exportacion",
        lambda db, periodo: [(evaluacion, resultado, colaborador)],
    )

    libro = load_workbook(io.BytesIO(reporte_service.exportar_evaluaciones_xlsx(None)))
    fecha = libro["Resultados"]["G5"].value
    assert fecha == datetime(2026, 10, 16, 8, 30)
    assert fecha.tzinfo is None
    libro.close()


def test_perfil_consolidado_e_historico(client, tokens):
    auth = headers(tokens["consulta"])
    r = client.get("/api/colaboradores?incluir_inactivos=true", headers=auth)
    col001 = None
    for item in r.json()["items"]:
        if item["codigo"] == "COL-001":
            col001 = item
    assert col001 is not None
    cid = col001["id"]

    r2 = client.get(f"/api/colaboradores/{cid}/perfil", headers=auth)
    assert r2.status_code == 200
    perfil = r2.json()
    assert perfil["colaborador"]["codigo"] == "COL-001"
    assert perfil["mediciones"]
    assert perfil["historial"], "El perfil debe incluir evaluaciones (RF-10)"

    r3 = client.get(f"/api/colaboradores/{cid}/historial", headers=auth)
    assert r3.status_code == 200
    periodos = {h["periodo"] for h in r3.json()}
    assert periodos  # Se conserva el historial de cada periodo (RF-17).


def test_consulta_asistida_ia_con_fallback(client, tokens):
    r = client.post(
        "/api/ia/consulta",
        json={"pregunta": "¿Qué colaboradores muestran mayor potencial este periodo?",
              "periodo": "2026-S1"},
        headers=headers(tokens["gestor"]),
    )
    assert r.status_code == 200
    cuerpo = r.json()
    assert cuerpo["origen"] == "local_determinista"
    assert cuerpo["respuesta"]
    assert cuerpo["advertencia"]  # sin credenciales Azure configuradas

    # El rol CONSULTA no tiene permiso de IA
    r2 = client.post(
        "/api/ia/consulta",
        json={"pregunta": "¿Quién es el mejor colaborador?"},
        headers=headers(tokens["consulta"]),
    )
    assert r2.status_code == 403


def test_regenerar_explicacion(client, tokens):
    r = client.get("/api/evaluaciones", params={"periodo": "2026-S1"},
                   headers=headers(tokens["consulta"]))
    ev = r.json()["items"][0]
    r2 = client.post(f"/api/evaluaciones/{ev['id']}/explicacion",
                     headers=headers(tokens["evaluador"]))
    assert r2.status_code == 200
    assert r2.json()["explicacion"]
    assert r2.json()["explicacion_origen"] == "local_determinista"


# Pruebas de las funciones de analisis.

def test_matriz_nueve_box(client, tokens):
    r = client.get("/api/reportes/matriz-9box", params={"periodo": "2026-S1"},
                   headers=headers(tokens["consulta"]))
    assert r.status_code == 200
    cuerpo = r.json()
    assert cuerpo["periodo"] == "2026-S1"
    assert cuerpo["total_evaluados"] > 0
    # Los niveles de las cajas son validos y sus totales coinciden con la poblacion.
    total_cajas = sum(len(v) for v in cuerpo["cajas"].values())
    assert total_cajas == cuerpo["total_evaluados"]
    for punto in (v for caja in cuerpo["cajas"].values() for v in caja):
        assert punto["nivel_desempeno"] in ("bajo", "medio", "alto")
        assert punto["nivel_potencial"] in ("bajo", "medio", "alto")
    # Existen las 9 etiquetas de la matriz
    assert len(cuerpo["etiquetas"]) == 9


def test_comparador_colaboradores(client, tokens):
    # Primero localizamos dos colaboradores con mediciones
    colabs = client.get("/api/colaboradores", headers=headers(tokens["consulta"])).json()["items"]
    ids = ",".join(str(c["id"]) for c in colabs[:3])

    r = client.get("/api/reportes/comparar", params={"ids": ids, "periodo": "2026-S1"},
                   headers=headers(tokens["consulta"]))
    assert r.status_code == 200
    cuerpo = r.json()
    assert 2 <= len(cuerpo["colaboradores"]) <= 3
    assert cuerpo["indicadores"], "Debe devolver indicadores comparados"
    for ind in cuerpo["indicadores"]:
        assert ind["radar"]
        assert ind["radar"].keys() == ind["valores"].keys()
        for valor in ind["normalizados"].values():
            assert 0 <= valor <= 100
        for colaborador_id, valor_radar in ind["radar"].items():
            valor = ind["valores"][colaborador_id]
            if ind["codigo"] == "ANIOS_EXPERIENCIA":
                esperado = (valor - 1) / 24 * 100
            elif ind["codigo"] == "ROTACION_EQUIPO":
                esperado = 100 - (valor - 2) / 28 * 100
            else:
                esperado = valor
            assert valor_radar == round(min(max(esperado, 0), 100), 1)

    # Se requiere al menos 2 colaboradores
    r2 = client.get("/api/reportes/comparar", params={"ids": str(colabs[0]["id"])},
                    headers=headers(tokens["consulta"]))
    assert r2.status_code == 422


def test_centro_de_alertas(client, tokens):
    r = client.get("/api/reportes/alertas", headers=headers(tokens["consulta"]))
    assert r.status_code == 200
    cuerpo = r.json()
    claves = {i["clave"] for i in cuerpo["items"]}
    assert "revisiones_pendientes" in claves
    assert "mediciones_pendientes" in claves
    assert "planes_en_riesgo" in claves
    assert cuerpo["total"] == sum(i["total"] for i in cuerpo["items"])


def test_carga_csv_validacion_y_registro(client, tokens):
    auth = headers(tokens["gestor"])

    # Plantilla descargable
    r0 = client.get("/api/indicadores/plantilla-csv", headers=auth)
    assert r0.status_code == 200
    assert "colaborador_codigo" in r0.text

    csv_contenido = (
        "colaborador_codigo;indicador_codigo;periodo;valor;validez\n"
        "COL-003;LIDERAZGO;2026-CSV;82.5;valida\n"
        "COL-999;LIDERAZGO;2026-CSV;90.0;valida\n"
        "COL-003;CODIGO_INEXISTENTE;2026-CSV;70.0;valida\n"
        "COL-003;COMUNICACION;2026-CSV;no-numero;valida\n"
        "COL-003;LIDERAZGO;2026-CSV;101;valida\n"
    )

    # En modo validacion no se guarda ningun dato (RN-01).
    r = client.post("/api/indicadores/carga-csv",
                    json={"contenido": csv_contenido, "validar_solo": True},
                    headers=auth)
    assert r.status_code == 200
    reporte = r.json()
    assert reporte["modo"] == "validacion"
    assert reporte["aceptadas"] == 1
    assert reporte["rechazadas"] == 4
    assert reporte["registradas"] == 0

    # En la carga normal solo se guarda la fila valida.
    r2 = client.post("/api/indicadores/carga-csv",
                     json={"contenido": csv_contenido},
                     headers=auth)
    assert r2.status_code == 200
    reporte2 = r2.json()
    assert reporte2["registradas"] == 1
    motivos = " | ".join(d["motivo"] for d in reporte2["detalle_rechazos"])
    assert "Colaborador inexistente" in motivos
    assert "Indicador inexistente" in motivos
    assert "Valor no numérico" in motivos
    assert "entre 0 y 100" in motivos

    # Se confirma que la medicion se guardo.
    colabs = client.get("/api/colaboradores?q=COL-003", headers=auth).json()["items"]
    r3 = client.get("/api/indicadores/mediciones",
                    params={"colaborador_id": colabs[0]["id"], "periodo": "2026-CSV"},
                    headers=auth)
    assert r3.status_code == 200
    assert len(r3.json()) == 1

    # Un encabezado incorrecto debe responder 422.
    r4 = client.post("/api/indicadores/carga-csv",
                     json={"contenido": "colaborador;indicador;periodo;valor\nCOL-003;X;P;1"},
                     headers=auth)
    assert r4.status_code == 422

    # CONSULTA no puede cargar CSV
    r5 = client.post("/api/indicadores/carga-csv",
                     json={"contenido": csv_contenido},
                     headers=headers(tokens["consulta"]))
    assert r5.status_code == 403


def test_plantilla_excel_formateada_y_carga_validacion(client, tokens):
    auth = headers(tokens["gestor"])
    plantilla = client.get("/api/indicadores/plantilla-xlsx", headers=auth)
    assert plantilla.status_code == 200
    assert "spreadsheetml.sheet" in plantilla.headers["content-type"]
    libro = load_workbook(io.BytesIO(plantilla.content))
    assert libro.sheetnames == ["Instrucciones", "Mediciones"]
    hoja = libro["Mediciones"]
    assert [hoja.cell(1, columna).value for columna in range(1, 6)] == [
        "colaborador_codigo", "indicador_codigo", "periodo", "valor", "validez",
    ]
    assert hoja["A1"].fill.fgColor.rgb.endswith("17365D")
    assert hoja.freeze_panes == "A2"
    for columna, valor in enumerate(
        ["COL-003", "LIDERAZGO", "2026-XLSX", 88.5, "valida"], start=1
    ):
        hoja.cell(2, columna, valor)
    for columna, valor in enumerate(
        ["COL-003", "COMUNICACION", "2026-XLSX", 101, "valida"], start=1
    ):
        hoja.cell(3, columna, valor)
    contenido = io.BytesIO()
    libro.save(contenido)
    libro.close()

    respuesta = client.post(
        "/api/indicadores/carga-xlsx",
        files={
            "archivo": (
                "mediciones.xlsx",
                contenido.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"validar_solo": "true"},
        headers=auth,
    )
    assert respuesta.status_code == 200, respuesta.text
    reporte = respuesta.json()
    assert reporte["modo"] == "validacion"
    assert reporte["total_filas"] == 2
    assert reporte["aceptadas"] == 1
    assert reporte["rechazadas"] == 1
    assert reporte["registradas"] == 0
    assert "entre 0 y 100" in reporte["detalle_rechazos"][0]["motivo"]


def test_dashboard_area_sin_puntuaciones(client, tokens):
    """Regresión: un área cuyos evaluados no tienen puntuación (sin mediciones
    válidas) no debe romper el dashboard con división entre cero."""
    auth_gestor = headers(tokens["gestor"])

    # Se agrega un colaborador en un area nueva y sin mediciones.
    r = client.post("/api/colaboradores", json={
        "codigo": "COL-SINDATOS", "nombre": "Colaborador Sin Datos",
        "area": "Área Sin Datos", "puesto": "Analista",
    }, headers=auth_gestor)
    assert r.status_code == 201

    # La evaluacion incluye a la persona, pero queda sin puntaje por falta de datos.
    r2 = client.post("/api/evaluaciones/ejecutar", json={"periodo": "2026-S1"},
                     headers=auth_gestor)
    assert r2.status_code == 200

    # El panel debe responder 200 y mostrar el area sin promedio.
    r3 = client.get("/api/reportes/dashboard", params={"periodo": "2026-S1"},
                    headers=headers(tokens["consulta"]))
    assert r3.status_code == 200
    areas = {a["area"]: a for a in r3.json()["promedios_por_area"]}
    assert "Área Sin Datos" in areas
    assert areas["Área Sin Datos"]["promedio"] is None
    assert areas["Área Sin Datos"]["evaluados"] == 0
