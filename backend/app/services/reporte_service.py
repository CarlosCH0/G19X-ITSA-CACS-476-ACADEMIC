# -*- coding: utf-8 -*-
"""Paneles, reportes y exportaciones (RF-12, RF-13, RF-14, RN-07)."""
import csv
import io
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (
    Clasificacion,
    Colaborador,
    Evaluacion,
    MedicionIndicador,
    ResultadoEvaluacion,
    Validez,
)
from app.services import plan_service

ZONA_HORARIA_CDMX = timezone(timedelta(hours=-6), name="CDMX")

# Etiquetas de la matriz 9-box: (desempeno, potencial) y su nombre.
ETIQUETAS_9BOX: dict[tuple[str, str], str] = {
    ("alto", "alto"): "Estrellas · sucesores listos",
    ("medio", "alto"): "Alto potencial emergente",
    ("bajo", "alto"): "Enigmas · potencial sin desempeño",
    ("alto", "medio"): "Alto desempeño",
    ("medio", "medio"): "Núcleo central",
    ("bajo", "medio"): "Contribuidores inconsistentes",
    ("alto", "bajo"): "Especialistas expertos",
    ("medio", "bajo"): "Contribuidores confiables",
    ("bajo", "bajo"): "En riesgo",
}


def _periodo_mas_reciente(db: Session) -> str | None:
    return db.scalar(select(func.max(Evaluacion.periodo)))


def dashboard(db: Session, periodo: str | None = None) -> dict:
    """Agregaciones para el panel de control."""
    periodo = periodo or _periodo_mas_reciente(db)

    total_activos = db.scalar(
        select(func.count()).select_from(Colaborador).where(Colaborador.estado.is_(True))
    ) or 0

    base = (
        select(Evaluacion, ResultadoEvaluacion, Colaborador)
        .join(ResultadoEvaluacion, ResultadoEvaluacion.evaluacion_id == Evaluacion.id)
        .join(Colaborador, Colaborador.id == Evaluacion.colaborador_id)
    )
    filas = db.execute(base).all()

    if periodo:
        filas = [f for f in filas if f[0].periodo == periodo]

    distribucion: dict[str, int] = {}
    por_area: dict[str, list[float]] = {}
    revision_estado: dict[str, int] = {"sin_revision": 0}
    for ev, res, colab in filas:
        clave = res.clasificacion.value if res.clasificacion else "sin_clasificacion"
        distribucion[clave] = distribucion.get(clave, 0) + 1
        por_area.setdefault(colab.area, [])
        if res.puntuacion is not None:
            por_area[colab.area].append(res.puntuacion)
        if res.revisiones:
            ultima = max(res.revisiones, key=lambda r: r.fecha)
            revision_estado[ultima.estado.value] = revision_estado.get(ultima.estado.value, 0) + 1
        else:
            revision_estado["sin_revision"] += 1

    promedios_area = [
        {
            "area": area,
            # Si el area no tiene puntuaciones, el promedio queda vacio
            # para evitar dividir entre cero.
            "promedio": round(sum(v) / len(v), 2) if v else None,
            "evaluados": len(v),
        }
        for area, v in sorted(por_area.items())
    ]

    hipo = [
        f for f in filas if f[1].clasificacion == Clasificacion.ALTO_POTENCIAL
    ]
    top = sorted(hipo, key=lambda f: (f[1].puntuacion or 0), reverse=True)[:10]
    top_hipo = [
        {
            "evaluacion_id": ev.id,
            "colaborador_id": colab.id,
            "codigo": colab.codigo,
            "nombre": colab.nombre,
            "area": colab.area,
            "puntuacion": res.puntuacion,
            "periodo": ev.periodo,
            "estado_revision": (
                max(res.revisiones, key=lambda r: r.fecha).estado.value
                if res.revisiones
                else "sin_revision"
            ),
        }
        for ev, res, colab in top
    ]

    # Se calcula el promedio de cada periodo para mostrar la tendencia.
    por_periodo: dict[str, list[float]] = {}
    for ev, res, _ in db.execute(base).all():
        if res.puntuacion is not None:
            por_periodo.setdefault(ev.periodo, []).append(res.puntuacion)
    evolucion = [
        {"periodo": p, "promedio": round(sum(v) / len(v), 2), "evaluados": len(v)}
        for p, v in sorted(por_periodo.items())
    ]

    return {
        "periodo": periodo,
        "generado_en": datetime.now(timezone.utc).isoformat(),
        "kpis": {
            "colaboradores_activos": total_activos,
            "evaluaciones_periodo": len(filas),
            "alto_potencial": distribucion.get(Clasificacion.ALTO_POTENCIAL.value, 0),
            "revisiones_pendientes": revision_estado.get("sin_revision", 0),
        },
        "distribucion_clasificacion": distribucion,
        "promedios_por_area": promedios_area,
        "revisiones": revision_estado,
        "top_alto_potencial": top_hipo,
        "evolucion_periodos": evolucion,
    }


COLUMNAS_EXPORTACION = [
    "evaluacion_id", "colaborador_codigo", "colaborador_nombre", "area", "puesto",
    "periodo", "fecha_evaluacion", "version_modelo", "configuracion_version",
    "estado_evaluacion", "puntuacion", "clasificacion", "cobertura",
    "factores_principales", "estado_revision", "explicacion_origen",
]


def _evaluaciones_exportacion(db: Session, periodo: str | None = None):
    consulta = (
        select(Evaluacion, ResultadoEvaluacion, Colaborador)
        .join(ResultadoEvaluacion, ResultadoEvaluacion.evaluacion_id == Evaluacion.id)
        .join(Colaborador, Colaborador.id == Evaluacion.colaborador_id)
        .order_by(Evaluacion.periodo, Colaborador.codigo)
    )
    filas = db.execute(consulta).all()
    if periodo:
        filas = [f for f in filas if f[0].periodo == periodo]
    return filas


def exportar_evaluaciones_csv(db: Session, periodo: str | None = None) -> str:
    """Exportación CSV de resultados (RF-14). Respeta permisos en el router."""
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(COLUMNAS_EXPORTACION)
    for ev, res, colab in _evaluaciones_exportacion(db, periodo):
        top = ", ".join(
            f"{f['nombre']}={f['contribucion']:.2f}" for f in (res.factores or [])[:3]
        )
        estado_rev = (
            max(res.revisiones, key=lambda r: r.fecha).estado.value if res.revisiones else ""
        )
        writer.writerow(
            [
                ev.id, colab.codigo, colab.nombre, colab.area, colab.puesto, ev.periodo,
                ev.fecha.strftime("%d/%m/%Y %H:%M:%S") if ev.fecha else "",
                ev.version_modelo,
                ev.configuracion.version if ev.configuracion else "",
                ev.estado.value, res.puntuacion if res.puntuacion is not None else "",
                res.clasificacion.value if res.clasificacion else "",
                f"{res.cobertura * 100:.0f}%", top, estado_rev, res.explicacion_origen or "",
            ]
        )
    return "\ufeff" + buffer.getvalue()


def exportar_evaluaciones_xlsx(db: Session, periodo: str | None = None) -> bytes:
    """Exportación Excel formateada de resultados (RF-14)."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resultados"
    hoja.sheet_view.showGridLines = False
    hoja.freeze_panes = "A5"
    hoja.sheet_properties.tabColor = "0F766E"

    azul = "17365D"
    azul_medio = "1F4E78"
    turquesa = "0F766E"
    fondo = "F3F7FB"
    blanco = "FFFFFF"
    texto = "24364B"
    borde = Side(style="thin", color="D9E2F0")
    ultima_columna = len(COLUMNAS_EXPORTACION)

    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_columna)
    titulo = hoja.cell(1, 1, "RESULTADOS DE EVALUACIONES")
    titulo.fill = PatternFill("solid", fgColor=azul)
    titulo.font = Font(name="Aptos Display", size=20, bold=True, color=blanco)
    titulo.alignment = Alignment(vertical="center", indent=1)
    for fila in hoja["A1:P1"]:
        for celda in fila:
            celda.fill = PatternFill("solid", fgColor=azul)
    hoja.row_dimensions[1].height = 38

    hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_columna)
    subtitulo = hoja.cell(
        2, 1,
        f"Sistema HiPo · Periodo: {periodo or 'todos'} · "
        f"Generado: {datetime.now(ZONA_HORARIA_CDMX):%d/%m/%Y %H:%M} hora CDMX",
    )
    subtitulo.font = Font(name="Aptos", size=10, italic=True, color=texto)
    subtitulo.alignment = Alignment(vertical="center", indent=1)
    hoja.row_dimensions[2].height = 24

    fila_encabezado = 4
    for columna, nombre in enumerate(COLUMNAS_EXPORTACION, start=1):
        celda = hoja.cell(fila_encabezado, columna, nombre)
        celda.fill = PatternFill("solid", fgColor=azul_medio)
        celda.font = Font(name="Aptos", size=10, bold=True, color=blanco)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = Border(bottom=Side(style="medium", color=turquesa))
    hoja.row_dimensions[fila_encabezado].height = 34

    anchos = [14, 20, 32, 24, 28, 16, 23, 18, 22, 22, 16, 22, 14, 48, 20, 24]
    for columna, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho

    for numero_fila, (ev, res, colab) in enumerate(
        _evaluaciones_exportacion(db, periodo), start=fila_encabezado + 1
    ):
        fecha_evaluacion = ev.fecha
        if fecha_evaluacion:
            if fecha_evaluacion.tzinfo is None:
                fecha_evaluacion = fecha_evaluacion.replace(tzinfo=timezone.utc)
            fecha_evaluacion = fecha_evaluacion.astimezone(ZONA_HORARIA_CDMX).replace(tzinfo=None)
        factores = ", ".join(
            f"{factor['nombre']}={factor['contribucion']:.2f}"
            for factor in (res.factores or [])[:3]
        )
        estado_revision = (
            max(res.revisiones, key=lambda revision: revision.fecha).estado.value
            if res.revisiones
            else ""
        )
        valores = [
            ev.id, colab.codigo, colab.nombre, colab.area, colab.puesto, ev.periodo,
            fecha_evaluacion, ev.version_modelo,
            ev.configuracion.version if ev.configuracion else "",
            ev.estado.value, res.puntuacion, res.clasificacion.value if res.clasificacion else "",
            res.cobertura, factores, estado_revision, res.explicacion_origen or "",
        ]
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(numero_fila, columna, valor)
            celda.fill = PatternFill(
                "solid", fgColor=blanco if numero_fila % 2 else fondo
            )
            celda.font = Font(name="Aptos", size=10, color=texto)
            celda.border = Border(bottom=borde)
            celda.alignment = Alignment(vertical="center", wrap_text=columna in (3, 13, 14))
            if isinstance(valor, str) and valor.startswith("="):
                celda.data_type = "s"
        hoja.cell(numero_fila, 7).number_format = "dd/mm/yyyy hh:mm:ss"
        hoja.cell(numero_fila, 11).number_format = "0.00"
        hoja.cell(numero_fila, 13).number_format = "0%"
        hoja.row_dimensions[numero_fila].height = 24

    ultima_fila = max(fila_encabezado, hoja.max_row)
    hoja.auto_filter.ref = f"A{fila_encabezado}:P{ultima_fila}"
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0

    salida = io.BytesIO()
    libro.save(salida)
    libro.close()
    return salida.getvalue()


# ------------------------------------------------------------------
# Matriz de sucesion 9-box: desempeno y potencial.
# ------------------------------------------------------------------
def _nivel_tercil(valor: float, cortes: tuple[float, float]) -> str:
    """Clasifica un valor en bajo/medio/alto según cortes de terciles."""
    if valor >= cortes[1]:
        return "alto"
    if valor >= cortes[0]:
        return "medio"
    return "bajo"


def _cortes_terciles(valores: list[float]) -> tuple[float, float]:
    """Cortes p33/p67 de la población (terciles aproximados)."""
    if not valores:
        return (0.0, 0.0)
    ordenados = sorted(valores)
    n = len(ordenados)
    p33 = ordenados[min(max(n // 3, 0), n - 1)]
    p67 = ordenados[min(max(2 * n // 3, 0), n - 1)]
    return (p33, p67)


def matriz_nueve_box(db: Session, periodo: str | None = None) -> dict:
    """Matriz de sucesión 9-box: desempeño (eje X) × potencial (eje Y).

    El desempeño se deriva de los factores tipo desempeño de cada resultado;
    el potencial es la puntuación global HiPo del periodo.
    """
    periodo = periodo or _periodo_mas_reciente(db)
    consulta = (
        select(Evaluacion, ResultadoEvaluacion, Colaborador)
        .join(ResultadoEvaluacion, ResultadoEvaluacion.evaluacion_id == Evaluacion.id)
        .join(Colaborador, Colaborador.id == Evaluacion.colaborador_id)
    )
    filas = db.execute(consulta).all()
    if periodo:
        filas = [f for f in filas if f[0].periodo == periodo]

    puntos = []
    for ev, res, colab in filas:
        if res.puntuacion is None:
            continue
        factores = res.factores or []
        desempenio_vals = [f["normalizado"] * 100 for f in factores if f.get("tipo") == "desempeno"]
        desempenio = (
            sum(desempenio_vals) / len(desempenio_vals) if desempenio_vals else res.puntuacion
        )
        puntos.append(
            {
                "evaluacion_id": ev.id,
                "colaborador_id": colab.id,
                "codigo": colab.codigo,
                "nombre": colab.nombre,
                "area": colab.area,
                "puesto": colab.puesto,
                "desempeno": round(desempenio, 1),
                "potencial": res.puntuacion,
                "clasificacion": res.clasificacion.value if res.clasificacion else None,
                "estado_revision": (
                    max(res.revisiones, key=lambda r: r.fecha).estado.value
                    if res.revisiones else "sin_revision"
                ),
            }
        )

    cortes_x = _cortes_terciles([p["desempeno"] for p in puntos])
    cortes_y = _cortes_terciles([p["potencial"] for p in puntos])
    cajas: dict[str, list[dict]] = {}
    for p in puntos:
        p["nivel_desempeno"] = _nivel_tercil(p["desempeno"], cortes_x)
        p["nivel_potencial"] = _nivel_tercil(p["potencial"], cortes_y)
        clave = f"{p['nivel_desempeno']}|{p['nivel_potencial']}"
        cajas.setdefault(clave, []).append(p)

    return {
        "periodo": periodo,
        "generado_en": datetime.now(timezone.utc).isoformat(),
        "cortes": {
            "desempeno": {"bajo": cortes_x[0], "alto": cortes_x[1]},
            "potencial": {"bajo": cortes_y[0], "alto": cortes_y[1]},
        },
        "etiquetas": {f"{d}|{p}": nombre for (d, p), nombre in ETIQUETAS_9BOX.items()},
        "cajas": cajas,
        "total_evaluados": len(puntos),
    }


# ------------------------------------------------------------------
# Comparador de colaboradores (RF-12: comparaciones)
# ------------------------------------------------------------------
def comparar_colaboradores(
    db: Session, ids: list[int], periodo: str | None = None
) -> dict:
    """Compara indicadores normalizados 0-100 entre 2..6 colaboradores."""
    ids = list(dict.fromkeys(ids))[:6]
    colaboradores = {c.id: c for c in db.scalars(select(Colaborador).where(Colaborador.id.in_(ids)))}
    if not colaboradores:
        return {"periodo": periodo, "colaboradores": [], "indicadores": []}

    consulta = (
        select(MedicionIndicador)
        .where(
            MedicionIndicador.colaborador_id.in_(ids),
            MedicionIndicador.validez == Validez.VALIDA,
        )
        .order_by(MedicionIndicador.colaborador_id)
    )
    if periodo:
        consulta = consulta.where(MedicionIndicador.periodo == periodo)
    mediciones = list(db.scalars(consulta))

    # Agrupar valores por indicador para normalizar min-max entre comparados
    por_indicador: dict[int, dict[int, float]] = {}
    nombres_indicador: dict[int, dict[str, str]] = {}
    for m in mediciones:
        por_indicador.setdefault(m.indicador_id, {})[m.colaborador_id] = m.valor
        nombres_indicador[m.indicador_id] = {
            "codigo": m.indicador.codigo,
            "nombre": m.indicador.nombre,
            "tipo": m.indicador.tipo.value,
            "direccion": m.indicador.direccion.value,
        }

    indicadores = []
    for indicador_id, valores in sorted(por_indicador.items()):
        mn, mx = min(valores.values()), max(valores.values())
        normalizados = {
            cid: (50.0 if mx == mn else round((v - mn) / (mx - mn) * 100, 1))
            for cid, v in valores.items()
        }
        meta = nombres_indicador[indicador_id]
        codigo = meta["codigo"]
        escala_radar = {
            "ANIOS_EXPERIENCIA": (1.0, 25.0),
            "ROTACION_EQUIPO": (2.0, 30.0),
        }.get(codigo)
        valores_radar: dict[int, float] = {}
        for cid, valor in valores.items():
            if escala_radar:
                minimo, maximo = escala_radar
                puntuacion_radar = (valor - minimo) / (maximo - minimo) * 100
            else:
                puntuacion_radar = valor
            if meta["direccion"] == "menor_mejor":
                puntuacion_radar = 100 - puntuacion_radar
            valores_radar[cid] = round(min(max(puntuacion_radar, 0), 100), 1)
        indicadores.append(
            {
                "indicador_id": indicador_id,
                "codigo": meta["codigo"],
                "nombre": meta["nombre"],
                "tipo": meta["tipo"],
                "min": mn,
                "max": mx,
                "valores": {str(cid): valores.get(cid) for cid in ids if cid in valores},
                "radar": {str(cid): valores_radar[cid] for cid in ids if cid in valores_radar},
                "normalizados": {
                    str(cid): normalizados.get(cid) for cid in ids if cid in normalizados
                },
            }
        )
    indicadores.sort(key=lambda i: (i["tipo"], i["codigo"]))

    # Puntaje y evaluacion del periodo para cada persona comparada.
    evaluaciones = {}
    consulta_ev = (
        select(Evaluacion, ResultadoEvaluacion)
        .join(ResultadoEvaluacion, ResultadoEvaluacion.evaluacion_id == Evaluacion.id)
        .where(Evaluacion.colaborador_id.in_(ids))
    )
    if periodo:
        consulta_ev = consulta_ev.where(Evaluacion.periodo == periodo)
    for ev, res in db.execute(consulta_ev):
        evaluaciones[ev.colaborador_id] = {
            "evaluacion_id": ev.id,
            "puntuacion": res.puntuacion,
            "clasificacion": res.clasificacion.value if res.clasificacion else None,
        }

    lista = [
        {
            "id": c.id,
            "codigo": c.codigo,
            "nombre": c.nombre,
            "area": c.area,
            "puesto": c.puesto,
            "evaluacion": evaluaciones.get(c.id),
        }
        for c in (colaboradores[i] for i in ids if i in colaboradores)
    ]
    return {
        "periodo": periodo or _periodo_mas_reciente(db),
        "colaboradores": lista,
        "indicadores": indicadores,
    }


# ------------------------------------------------------------------
# Centro de alertas (calidad de datos y pendientes)
# ------------------------------------------------------------------
def alertas(db: Session) -> dict:
    """Resumen de pendientes operativos: revisiones, validaciones y planes."""
    filas = db.execute(
        select(Evaluacion, ResultadoEvaluacion)
        .join(ResultadoEvaluacion, ResultadoEvaluacion.evaluacion_id == Evaluacion.id)
    ).all()
    periodo = _periodo_mas_reciente(db)

    sin_revision = [
        (ev, res)
        for ev, res in filas
        if (periodo is None or ev.periodo == periodo) and not res.revisiones
    ]
    sin_explicacion = [
        (ev, res)
        for ev, res in filas
        if (periodo is None or ev.periodo == periodo) and not res.explicacion
    ]

    pendientes = db.scalar(
        select(func.count()).select_from(MedicionIndicador).where(
            MedicionIndicador.validez == Validez.PENDIENTE
        )
    ) or 0
    invalidas = db.scalar(
        select(func.count()).select_from(MedicionIndicador).where(
            MedicionIndicador.validez == Validez.INVALIDA
        )
    ) or 0

    stats_planes = plan_service.estadisticas(db)

    detalle_revisiones = [
        {
            "evaluacion_id": ev.id,
            "colaborador_id": ev.colaborador_id,
            "colaborador": ev.colaborador.nombre if ev.colaborador else "",
            "periodo": ev.periodo,
            "puntuacion": res.puntuacion,
            "clasificacion": res.clasificacion.value if res.clasificacion else None,
        }
        for ev, res in sorted(sin_revision, key=lambda f: f[1].puntuacion or 0, reverse=True)[:10]
    ]

    items = [
        {
            "clave": "revisiones_pendientes",
            "titulo": "Resultados sin revisión humana",
            "detalle": f"{len(sin_revision)} resultados del periodo {periodo or '—'} esperan confirmación (RN-06).",
            "total": len(sin_revision),
        },
        {
            "clave": "mediciones_pendientes",
            "titulo": "Mediciones pendientes de validar",
            "detalle": f"{pendientes} mediciones en estado PENDIENTE no participan en evaluaciones (RN-01).",
            "total": pendientes,
        },
        {
            "clave": "mediciones_invalidas",
            "titulo": "Mediciones marcadas inválidas",
            "detalle": f"{invalidas} mediciones invalidadas requieren corrección o recarga.",
            "total": invalidas,
        },
        {
            "clave": "planes_en_riesgo",
            "titulo": "Planes de desarrollo en riesgo",
            "detalle": f"{stats_planes['en_riesgo']} planes con acciones vencidas o fecha superada.",
            "total": stats_planes["en_riesgo"],
        },
        {
            "clave": "planes_vencidos",
            "titulo": "Planes vencidos sin cerrar",
            "detalle": f"{stats_planes['vencidos']} planes abiertos con fecha fin pasada.",
            "total": stats_planes["vencidos"],
        },
        {
            "clave": "sin_explicacion",
            "titulo": "Resultados sin explicación registrada",
            "detalle": f"{len(sin_explicacion)} resultados aún no generan explicación (RF-09).",
            "total": len(sin_explicacion),
        },
    ]
    return {
        "periodo": periodo,
        "generado_en": datetime.now(timezone.utc).isoformat(),
        "total": sum(i["total"] for i in items),
        "items": items,
        "revisiones_destacadas": detalle_revisiones,
    }
